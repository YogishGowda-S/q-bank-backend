from flask import Flask, request, jsonify
from flask_cors import CORS
import os
import re
import docx
import pdfplumber
from groq import Groq
import psycopg2
from psycopg2.extras import RealDictCursor
import time
import traceback
import hashlib

app = Flask(__name__)
CORS(app)
app.config['MAX_CONTENT_LENGTH'] = 10 * 1024 * 1024

# Get from environment
DATABASE_URL = os.getenv("DATABASE_URL")
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

try:
    groq_client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None
except:
    groq_client = None

def get_db():
    try:
        conn = psycopg2.connect(DATABASE_URL)
        return conn
    except Exception as e:
        print(f"DB connection error: {e}")
        return None

def init_db():
    try:
        conn = get_db()
        if not conn:
            return
        cursor = conn.cursor()
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS users (
                id SERIAL PRIMARY KEY,
                email TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS subjects (
                id SERIAL PRIMARY KEY,
                user_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        ''')
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS questions (
                id SERIAL PRIMARY KEY,
                subject_id INTEGER NOT NULL,
                question_text TEXT NOT NULL,
                marks INTEGER,
                FOREIGN KEY (subject_id) REFERENCES subjects(id)
            )
        ''')
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS answers (
                id SERIAL PRIMARY KEY,
                question_id INTEGER NOT NULL,
                answer_text TEXT,
                word_count INTEGER,
                FOREIGN KEY (question_id) REFERENCES questions(id)
            )
        ''')
        
        conn.commit()
        cursor.close()
        conn.close()
        print("Database initialized")
    except Exception as e:
        print(f"DB init error: {e}")

init_db()

def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def parse_questions(docx_paths):
    questions = []
    seen = set()
    skip_keywords = [
        'subject name', 'subject code', 'question bank',
        'department', 'semester', 'module', 'unit', 'section',
        'part a', 'part b', 'sl no', 'sl.no', 'course code',
        'course name', 'faculty', 'year', 'batch', 'college',
        'ia1', 'ia2', 'ia3', 'internal assessment'
    ]

    for docx_path in docx_paths:
        try:
            doc = docx.Document(docx_path)
            para_count = 0
            for para in doc.paragraphs:
                if para_count > 100:
                    break
                text = para.text.strip()
                if not text or len(text) < 15 or text in seen or any(kw in text.lower() for kw in skip_keywords) or len(text.split()) < 5:
                    continue
                seen.add(text)
                match = re.search(r'\b(\d+)\s*[Mm]arks?\b|\[(\d+)\]|\((\d+)\)|\b(\d+)[Mm]\b', text)
                marks = int(next((m for m in match.groups() if m), 0)) if match else 0
                questions.append({"text": text, "marks": marks})
                para_count += 1

            table_count = 0
            for table in doc.tables:
                if table_count > 5:
                    break
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if not text or len(text) < 15 or text in seen or any(kw in text.lower() for kw in skip_keywords) or len(text.split()) < 5:
                            continue
                        seen.add(text)
                        match = re.search(r'\b(\d+)\s*[Mm]arks?\b|\[(\d+)\]|\((\d+)\)|\b(\d+)[Mm]\b', text)
                        marks = int(next((m for m in match.groups() if m), 0)) if match else 0
                        questions.append({"text": text, "marks": marks})
                table_count += 1
        except Exception as e:
            print(f"Parse error: {e}")

    return questions[:30]

def parse_notes(pdf_paths):
    all_text = ""
    for pdf_path in pdf_paths:
        try:
            with pdfplumber.open(pdf_path) as pdf:
                page_count = 0
                for page in pdf.pages:
                    if page_count > 10:
                        break
                    text = page.extract_text()
                    if text:
                        all_text += text + "\n\n"
                    page_count += 1
        except Exception as e:
            print(f"PDF parse error: {e}")
    
    return all_text[:40000]

def clean_answer(raw):
    raw = re.sub(r'<think>.*?</think>', '', raw, flags=re.DOTALL)
    raw = re.sub(r'<thinking>.*?</thinking>', '', raw, flags=re.DOTALL)
    raw = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', raw)
    raw = re.sub(r'\*(.*?)\*', r'<em>\1</em>', raw)
    raw = raw.replace('\n\n', '<br><br>').replace('\n', '<br>')
    return raw.strip()[:2500]

def generate_one_answer(question, notes_text):
    if not groq_client:
        marks = question["marks"] if question["marks"] > 0 else 10
        return f"<strong>Answer ({marks} marks):</strong><br>API not configured"

    notes_snippet = notes_text[:800] if len(notes_text) > 800 else notes_text
    marks = question["marks"] if question["marks"] > 0 else 10

    if marks <= 2:
        length_guide = "Write 2-3 lines. Be concise."
    elif marks <= 5:
        length_guide = "Write 5-8 lines. Include key points."
    else:
        length_guide = "Write 10-20 lines. Include detailed explanation."

    prompt = f"""Answer this exam question for engineering students. Write ONLY the answer, no preamble.

Question: {question['text']}
Marks: {marks}
{length_guide}

Start immediately with the answer."""

    for attempt in range(3):
        try:
            response = groq_client.chat.completions.create(
                model="llama-3.1-8b-instant",
                messages=[
                    {
                        "role": "system",
                        "content": "You are an exam answer writer. Write clear, concise answers for engineering exams. No thinking process, no preamble. Start immediately."
                    },
                    {
                        "role": "user",
                        "content": prompt
                    }
                ],
                max_tokens=1000,
                temperature=0.3,
                timeout=60
            )
            raw = response.choices[0].message.content.strip()
            return clean_answer(raw)
        except Exception as e:
            error_str = str(e).lower()
            if "429" in str(e) or "rate" in error_str:
                if attempt < 2:
                    time.sleep(30)
                    continue
                return f"<strong>Answer ({marks} marks):</strong><br>Rate limited. Please try again."
            elif "timeout" in error_str:
                if attempt < 2:
                    time.sleep(20)
                    continue
                return f"<strong>Answer ({marks} marks):</strong><br>Response timeout."
            else:
                if attempt < 2:
                    time.sleep(10)
                    continue
    
    return f"<strong>Answer ({marks} marks):</strong><br>Could not generate answer."

def generate_all_answers(questions, notes_text):
    all_answers = []
    for i, question in enumerate(questions):
        try:
            print(f"Q{i+1}/{len(questions)}")
            answer = generate_one_answer(question, notes_text)
            all_answers.append(answer)
            time.sleep(1)
        except Exception as e:
            print(f"Error Q{i+1}: {e}")
            marks = question.get("marks", 10)
            all_answers.append(f"<strong>Answer ({marks} marks):</strong><br>Error")
    
    return all_answers

# API ENDPOINTS

@app.route("/api/signup", methods=["POST"])
def signup():
    try:
        data = request.json
        email = data.get("email", "").strip()
        password = data.get("password", "").strip()
        
        if not email or not password:
            return jsonify({"error": "Email and password required"}), 400
        
        password_hash = hash_password(password)
        
        conn = get_db()
        if not conn:
            return jsonify({"error": "Database error"}), 500
        
        cursor = conn.cursor()
        cursor.execute("INSERT INTO users (email, password_hash) VALUES (%s, %s)", (email, password_hash))
        conn.commit()
        cursor.close()
        conn.close()
        
        return jsonify({"message": "Signup successful"}), 201
    except Exception as e:
        if "duplicate" in str(e).lower():
            return jsonify({"error": "Email already exists"}), 400
        return jsonify({"error": str(e)[:100]}), 500

@app.route("/api/login", methods=["POST"])
def login():
    try:
        data = request.json
        email = data.get("email", "").strip()
        password = data.get("password", "").strip()
        
        if not email or not password:
            return jsonify({"error": "Email and password required"}), 400
        
        password_hash = hash_password(password)
        
        conn = get_db()
        if not conn:
            return jsonify({"error": "Database error"}), 500
        
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute("SELECT id FROM users WHERE email=%s AND password_hash=%s", (email, password_hash))
        user = cursor.fetchone()
        cursor.close()
        conn.close()
        
        if user:
            return jsonify({"user_id": user['id'], "message": "Login successful"}), 200
        else:
            return jsonify({"error": "Invalid credentials"}), 401
    except Exception as e:
        return jsonify({"error": str(e)[:100]}), 500

@app.route("/api/subjects", methods=["GET"])
def get_subjects():
    try:
        user_id = request.args.get("user_id")
        if not user_id:
            return jsonify({"error": "user_id required"}), 400
        
        conn = get_db()
        if not conn:
            return jsonify({"error": "Database error"}), 500
        
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute("SELECT id, name FROM subjects WHERE user_id=%s ORDER BY created_at DESC", (user_id,))
        subjects = cursor.fetchall()
        cursor.close()
        conn.close()
        
        return jsonify({"subjects": subjects}), 200
    except Exception as e:
        return jsonify({"error": str(e)[:100]}), 500

@app.route("/api/subjects", methods=["POST"])
def create_subject():
    try:
        data = request.json
        user_id = data.get("user_id")
        name = data.get("name", "").strip()
        
        if not user_id or not name:
            return jsonify({"error": "user_id and name required"}), 400
        
        conn = get_db()
        if not conn:
            return jsonify({"error": "Database error"}), 500
        
        cursor = conn.cursor()
        cursor.execute("INSERT INTO subjects (user_id, name) VALUES (%s, %s) RETURNING id", (user_id, name))
        subject_id = cursor.fetchone()[0]
        conn.commit()
        cursor.close()
        conn.close()
        
        return jsonify({"subject_id": subject_id, "message": "Subject created"}), 201
    except Exception as e:
        return jsonify({"error": str(e)[:100]}), 500

@app.route("/api/subjects/<int:subject_id>", methods=["GET"])
def get_subject(subject_id):
    try:
        conn = get_db()
        if not conn:
            return jsonify({"error": "Database error"}), 500
        
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute("SELECT id, name FROM subjects WHERE id=%s", (subject_id,))
        subject = cursor.fetchone()
        
        cursor.execute("SELECT q.id, q.question_text, q.marks, a.answer_text FROM questions q LEFT JOIN answers a ON q.id = a.question_id WHERE q.subject_id=%s", (subject_id,))
        qa_pairs = cursor.fetchall()
        cursor.close()
        conn.close()
        
        if not subject:
            return jsonify({"error": "Subject not found"}), 404
        
        return jsonify({"subject": subject, "qa_pairs": qa_pairs}), 200
    except Exception as e:
        return jsonify({"error": str(e)[:100]}), 500

@app.route("/api/upload/<int:subject_id>", methods=["POST"])
def upload(subject_id):
    try:
        qbank_files = request.files.getlist("qbank")
        notes_files = request.files.getlist("notes")
        
        if not qbank_files or not qbank_files[0].filename:
            return jsonify({"error": "Please upload a question bank file"}), 400
        
        upload_folder = f"uploads/{subject_id}"
        os.makedirs(upload_folder, exist_ok=True)
        
        qbank_paths = []
        for i, f in enumerate(qbank_files[:2]):
            if f and f.filename:
                path = os.path.join(upload_folder, f"qbank_{i}.docx")
                f.save(path)
                qbank_paths.append(path)
        
        notes_paths = []
        for i, f in enumerate(notes_files[:2]):
            if f and f.filename:
                path = os.path.join(upload_folder, f"notes_{i}.pdf")
                f.save(path)
                notes_paths.append(path)
        
        questions = parse_questions(qbank_paths)
        notes_text = parse_notes(notes_paths)
        
        if not questions:
            return jsonify({"error": "No questions found in files"}), 400
        
        conn = get_db()
        if not conn:
            return jsonify({"error": "Database error"}), 500
        
        cursor = conn.cursor()
        cursor.execute("DELETE FROM answers WHERE question_id IN (SELECT id FROM questions WHERE subject_id=%s)", (subject_id,))
        cursor.execute("DELETE FROM questions WHERE subject_id=%s", (subject_id,))
        conn.commit()
        
        print(f"Generating answers for {len(questions)} questions...")
        answers = generate_all_answers(questions, notes_text)
        
        for question, answer in zip(questions, answers):
            cursor.execute("INSERT INTO questions (subject_id, question_text, marks) VALUES (%s, %s, %s) RETURNING id", 
                         (subject_id, question["text"][:500], question["marks"]))
            question_id = cursor.fetchone()[0]
            cursor.execute("INSERT INTO answers (question_id, answer_text, word_count) VALUES (%s, %s, %s)", 
                         (question_id, answer[:3000], len(answer.split())))
        
        conn.commit()
        cursor.close()
        conn.close()
        
        return jsonify({"message": "Answers generated successfully", "count": len(questions)}), 200
    except Exception as e:
        print(f"Upload error: {e}")
        traceback.print_exc()
        return jsonify({"error": str(e)[:200]}), 500

@app.route("/api/download/<int:subject_id>", methods=["GET"])
def download(subject_id):
    try:
        conn = get_db()
        if not conn:
            return jsonify({"error": "Database error"}), 500
        
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute("SELECT q.question_text, q.marks, a.answer_text FROM questions q LEFT JOIN answers a ON q.id = a.question_id WHERE q.subject_id=%s", (subject_id,))
        rows = cursor.fetchall()
        
        cursor.execute("SELECT name FROM subjects WHERE id=%s", (subject_id,))
        subject = cursor.fetchone()
        cursor.close()
        conn.close()
        
        if not rows:
            return jsonify({"error": "No answers to download"}), 404
        
        # Return data for PDF generation
        return jsonify({
            "subject_name": subject['name'],
            "qa_pairs": rows
        }), 200
    except Exception as e:
        return jsonify({"error": str(e)[:100]}), 500

@app.route("/health", methods=["GET"])
def health():
    return jsonify({"status": "ok"}), 200

if __name__ == "__main__":
    app.run(debug=True)
